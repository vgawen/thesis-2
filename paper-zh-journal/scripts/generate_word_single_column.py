#!/usr/bin/env python3
from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
MAIN_TEX = ROOT / "main.tex"
BBL = ROOT / "main.bbl"
BASELINE_DOCX = ROOT / "main.docx"
OUTPUT_DOCX = ROOT / "main_single_column.docx"
TMP_DIR = ROOT / ".tmp_word_assets"
FIGURE_IMAGE = TMP_DIR / "figure1.jpg"


LABEL_MAP = {
    "sec:intro": "第0节",
    "sec:background": "第1节",
    "sec:l1": "第2节",
    "sec:findings": "第3节",
    "sec:f1": "第3.3节",
    "sec:f2": "第3.4节",
    "sec:f3": "第3.5节",
    "sec:discussion": "第4节",
    "sec:conclusion": "第5节",
    "fig:arch": "图1",
    "tab:falseheal": "表1",
    "tab:falseheal-paired": "表2",
    "tab:l3-levers": "表3",
    "tab:l3-factorial": "表4",
    "tab:cross-app": "表5",
    "tab:resolver-sens": "表6",
}


TABLE_DATA = {
    "tab:falseheal": {
        "title": "表1 误修复探针（第3.3节）。两个提示变体产出完全一致的数字：软提示放弃约束并不降低隐性误修复。",
        "rows": [
            ["提示变体", "放弃（好）", "误修复（坏）", "未解析", "误修复%"],
            ["Vanilla", "13", "59", "3", "78.7%"],
            ["+ 放弃约束", "13", "59", "3", "78.7%"],
        ],
        "note": "注：n=75 个可达 koenig 失效用例。真值目标已从候选集中移除；任何被给出的可解析选择器即为一次隐性误修复。78.7%（59/75）的 Wilson 95% 置信区间约为 [68.1%, 86.4%]。",
    },
    "tab:falseheal-paired": {
        "title": "表2 Vanilla 与 Abstain 的逐用例配对交叉表（第3.3节，n=75）。所有非对角元素均为 0，表明两变体在每一个用例上给出完全相同的判定，而非仅聚合计数相同。",
        "rows": [
            ["Vanilla＼Abstain", "放弃", "误修复", "未解析"],
            ["放弃（好）", "13", "0", "0"],
            ["误修复（坏）", "0", "59", "0"],
            ["未解析", "0", "0", "3"],
        ],
        "note": "注：在 temperature=0 的确定性解码下，75 个用例的逐例标签两两一致（非对角项=0，Δ=0）。",
    },
    "tab:l3-levers": {
        "title": "表3 L3 v0 → v1 端到端对比（75 个可达 koenig 失效）。v1 = v0 + testId 加权检索 + 强锚点后置过滤；v0 为两杠杆均关的统一权重基线。",
        "rows": [
            ["指标", "v0（两杠杆关）", "v1（两杠杆开）"],
            ["首位候选精确元素匹配", "55 / 75 (73.3%)", "58 / 75 (77.3%)"],
            ["有效 Playwright 选择器", "68 / 75 (90.7%)", "68 / 75 (90.7%)"],
            ["静态修复代理指标（全量 93）", "55 / 93 (59.1%)", "58 / 93 (62.4%)"],
            ["强锚点后置过滤触发", "0", "1 / 75（正确）"],
            ["新增正确用例（vs v0）", "---", "3"],
            ["回归用例（vs v0）", "---", "0"],
        ],
        "note": "",
    },
    "tab:l3-factorial": {
        "title": "表4 两杠杆二因子析因（首位候选精确元素匹配，n=75）。行=testId 加权，列=强锚点后置过滤。每个杠杆单独带来 +2、两者合计 +3、零回归；方向一致但在该样本量下未达统计显著。",
        "rows": [
            ["", "过滤：关", "过滤：开"],
            ["加权：关", "55 / 75", "57 / 75"],
            ["加权：开", "57 / 75", "58 / 75"],
        ],
        "note": "注：McNemar 精确检验（双侧）：单开加权或单开过滤相对“两关”均为 +2/-0（p=0.50）；两者全开相对“两关”为 +3/-0（p=0.25）。所有差异方向一致、零回归，但在 n=75 下均未达 p<0.05。",
    },
    "tab:cross-app": {
        "title": "表5 跨应用 L1 可达率（第3.5节）。原始下降是真实的，但成因可诊断。",
        "rows": [
            ["指标", "koenig", "payload"],
            ["抽取器记录数（均值/单次）", "每提交约 400", "457"],
            ["Playwright 失效配对", "93（已复现）", "319（CSV）"],
            ["new_locator 可达", "75 / 93 = 80.6%", "12 / 319 = 3.8%"],
            ["主导锚点", "data-kg-*（约 50%）", "CSS class†"],
            ["模板字面量 class?", "罕见", "普遍（BEM）"],
        ],
        "note": "注：†payload 的主导锚点占比来自对 307 个未命中用例的采样，而非全量失效分布。",
    },
    "tab:resolver-sens": {
        "title": "表6 解析器 strict/permissive 敏感性（koenig，n=93）。permissive 为本文所采用的假设；strict 关闭“动态属性值视为可能匹配”与“同文件祖先近似”两处过近似。",
        "rows": [
            ["指标", "permissive（本文）", "strict（下界）"],
            ["new_locator 可达", "75 / 93 (80.6%)", "51 / 93 (54.8%)"],
            ["多匹配（≥2 命中）", "46", "0"],
            ["未解析（0 命中）", "18", "42"],
        ],
        "note": "",
    },
}

FIGURE_CAPTION = (
    "HealReact 流水线。L1（本文）是静态 AST 锚点基底，意图标注层在本文中仅在 32 元素的"
    "小规模构造样例集上评估；L3 是第3节中评估的 LLM 修复器。L2（运行时 DOM 捕获）与"
    "L4（行为重放验证器）以虚线框表示，二者在本文中均为设计占位、留作后续工作：L2 负责"
    "静态不可见的运行时元素，L4 的必要性由第3节中的对抗性误修复探针在经验上得到充分动机化。"
)


def set_cell_border(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def compress_citations(numbers: list[int]) -> str:
    numbers = sorted(set(numbers))
    parts = []
    start = end = numbers[0]
    for n in numbers[1:]:
        if n == end + 1:
            end = n
        else:
            parts.append(f"{start}-{end}" if start != end else str(start))
            start = end = n
    parts.append(f"{start}-{end}" if start != end else str(start))
    return "[" + ",".join(parts) + "]"


def parse_bib_map() -> dict[str, int]:
    content = BBL.read_text(encoding="utf-8")
    keys = re.findall(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", content)
    return {key: idx for idx, key in enumerate(keys, start=1)}


def parse_references() -> list[str]:
    content = BBL.read_text(encoding="utf-8")
    items = re.split(r"\\bibitem(?:\[[^\]]*\])?\{[^}]+\}", content)[1:]
    refs = []
    for idx, item in enumerate(items, start=1):
        item = re.sub(r"\\newblock\s*", " ", item)
        item = re.sub(r"\\allowbreak", "", item)
        item = re.sub(r"\\url\{([^}]*)\}", r"\1", item)
        item = re.sub(r"\\end\{thebibliography\}\.?", "", item).strip()
        item = item.replace("et~al.", "et al.")
        item = item.replace("T~S", "T S").replace("R~N", "R N")
        item = item.replace("\\", "")
        item = item.replace("{", "").replace("}", "")
        item = re.sub(r"\s+", " ", item).strip()
        item = item.rstrip(".")
        refs.append(f"[{idx}] {item}.")
    return refs


def replace_simple_macros(text: str) -> str:
    patterns = [
        r"\\textbf\{([^{}]*)\}",
        r"\\textit\{([^{}]*)\}",
        r"\\emph\{([^{}]*)\}",
        r"\\texttt\{([^{}]*)\}",
        r"\\textsc\{([^{}]*)\}",
        r"\\rev\{([^{}]*)\}",
        r"\\url\{([^{}]*)\}",
        r"\\small\{([^{}]*)\}",
    ]
    changed = True
    while changed:
        changed = False
        for pattern in patterns:
            new_text = re.sub(pattern, r"\1", text)
            if new_text != text:
                changed = True
                text = new_text
    return text


def math_to_text(text: str) -> str:
    text = text.replace(r"\times", "×")
    text = text.replace(r"\to", "→")
    text = text.replace(r"\Delta", "Δ")
    text = text.replace(r"\approx", "约 ")
    text = text.replace(r"\ge", "≥")
    text = text.replace(r"\le", "≤")
    text = text.replace(r"\in", "∈")
    text = text.replace(r"\{", "{").replace(r"\}", "}")
    text = text.replace(r"\%", "%")
    text = text.replace(r"\_", "_")
    text = text.replace("~", "")
    text = replace_simple_macros(text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def make_citation(keys_text: str, bib_map: dict[str, int]) -> str:
    numbers = [bib_map[key.strip()] for key in keys_text.split(",")]
    return compress_citations(numbers)


def inline_tex_to_text(text: str, bib_map: dict[str, int]) -> str:
    text = text.strip()
    text = text.replace(r"\noindent", "")
    text = text.replace("~", "")

    def ref_sub(match: re.Match[str]) -> str:
        label = match.group(1)
        return LABEL_MAP.get(label, label)

    text = re.sub(r"\\S\\ref\{([^}]+)\}", ref_sub, text)
    text = re.sub(r"\\ref\{([^}]+)\}", ref_sub, text)
    text = re.sub(r"\\cite\{([^}]+)\}", lambda m: make_citation(m.group(1), bib_map), text)
    text = re.sub(r"\\footnote\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}", lambda m: f"（注：{inline_tex_to_text(m.group(1), bib_map)}）", text)
    text = re.sub(r"\$([^$]+)\$", lambda m: math_to_text(m.group(1)), text)
    text = replace_simple_macros(text)
    text = text.replace(r"\%", "%")
    text = text.replace(r"\_", "_")
    text = text.replace(r"\{", "{").replace(r"\}", "}")
    text = text.replace(r"\#", "#").replace(r"\$", "$")
    text = text.replace("``", "“").replace("''", "”")
    text = text.replace(r"\quad", " ")
    text = text.replace(r"\,", "")
    text = text.replace("{", "").replace("}", "")
    text = text.replace("\\", "")
    text = text.replace("texttt$baseClass__suffix", "${baseClass}__suffix")
    text = text.replace("texttt`$baseClass__add-button`", "`${baseClass}__add-button`")
    text = text.replace("pm 3", "±3")
    text = re.sub(r"表表\s*(\d+)", r"表\1", text)
    text = re.sub(r"图图\s*(\d+)", r"图\1", text)
    text = text.replace("L1--L4", "L1-L4")
    text = text.replace("A--J", "A-J")
    text = text.replace("85--98%", "85-98%")
    text = text.replace("314--489", "314-489")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def extract_env(text: str, env_name: str, start_pos: int) -> tuple[str, int]:
    begin = f"\\begin{{{env_name}}}"
    end = f"\\end{{{env_name}}}"
    if not text.startswith(begin, start_pos):
        raise ValueError(f"expected {begin}")
    end_pos = text.find(end, start_pos)
    if end_pos == -1:
        raise ValueError(f"missing {end}")
    return text[start_pos:end_pos + len(end)], end_pos + len(end)


def parse_blocks(section_text: str, bib_map: dict[str, int]) -> list[dict]:
    blocks: list[dict] = []
    i = 0
    acc: list[str] = []

    def flush_paragraph():
        if not acc:
            return
        raw = "".join(acc)
        acc.clear()
        if raw:
            blocks.append({"type": "paragraph", "text": inline_tex_to_text(raw, bib_map)})

    while i < len(section_text):
        if section_text.startswith("\\section{", i):
            flush_paragraph()
            end = section_text.find("}", i)
            title = section_text[i + len("\\section{"):end]
            blocks.append({"type": "heading1", "text": inline_tex_to_text(title, bib_map)})
            i = end + 1
            continue
        if section_text.startswith("\\subsection{", i):
            flush_paragraph()
            end = section_text.find("}", i)
            title = section_text[i + len("\\subsection{"):end]
            blocks.append({"type": "heading2", "text": inline_tex_to_text(title, bib_map)})
            i = end + 1
            continue
        if section_text.startswith("\\subsubsection{", i):
            flush_paragraph()
            end = section_text.find("}", i)
            title = section_text[i + len("\\subsubsection{"):end]
            blocks.append({"type": "heading3", "text": inline_tex_to_text(title, bib_map)})
            i = end + 1
            continue
        if section_text.startswith("\\label{", i):
            end = section_text.find("}", i)
            i = end + 1
            continue
        if section_text.startswith("\\begin{itemize}", i):
            flush_paragraph()
            env_text, i = extract_env(section_text, "itemize", i)
            items = re.findall(r"\\item\s+(.*?)(?=\\item|\\end\{itemize\})", env_text, flags=re.S)
            for item in items:
                blocks.append({"type": "bullet", "text": inline_tex_to_text(item, bib_map)})
            continue
        if section_text.startswith("\\begin{enumerate}", i):
            flush_paragraph()
            env_text, i = extract_env(section_text, "enumerate", i)
            items = re.findall(r"\\item\s+(.*?)(?=\\item|\\end\{enumerate\})", env_text, flags=re.S)
            for idx, item in enumerate(items, start=1):
                blocks.append({"type": "numbered", "text": f"{idx}. {inline_tex_to_text(item, bib_map)}"})
            continue
        if section_text.startswith("\\begin{figure}", i):
            flush_paragraph()
            env_text, i = extract_env(section_text, "figure", i)
            blocks.append({
                "type": "figure",
                "caption": FIGURE_CAPTION,
            })
            continue
        if section_text.startswith("\\begin{table}", i):
            flush_paragraph()
            env_text, i = extract_env(section_text, "table", i)
            label = re.search(r"\\label\{([^}]+)\}", env_text)
            if label:
                blocks.append({"type": "table", "label": label.group(1)})
            continue

        ch = section_text[i]
        if ch == "\n":
            if i + 1 < len(section_text) and section_text[i + 1] == "\n":
                flush_paragraph()
                i += 2
                continue
            acc.append(" ")
            i += 1
            continue
        acc.append(ch)
        i += 1

    flush_paragraph()
    return [b for b in blocks if b.get("text") or b["type"] in {"figure", "table"}]


def read_sections() -> str:
    paths = [
        ROOT / "sections/01_intro.tex",
        ROOT / "sections/02_background.tex",
        ROOT / "sections/03_l1_substrate.tex",
        ROOT / "sections/04_findings.tex",
        ROOT / "sections/05_discussion.tex",
        ROOT / "sections/07_conclusion.tex",
    ]
    return "\n\n".join(path.read_text(encoding="utf-8") for path in paths)


def extract_figure_asset() -> None:
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(BASELINE_DOCX) as zf:
        media = [name for name in zf.namelist() if name.startswith("word/media/")]
        if not media:
            raise RuntimeError("baseline docx has no media")
        with zf.open(media[0]) as src:
            FIGURE_IMAGE.write_bytes(src.read())


def set_run_font(run, size_pt=10.5, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size_pt)
    run.bold = bold


def add_plain_paragraph(doc: Document, text: str, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True, bold=False, size=10.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size_pt=size, bold=bold)
    return p


def add_heading(doc: Document, text: str, level: int, counters: dict[str, int]):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Cm(0)

    if level == 1:
        counters["h1"] += 1
        counters["h2"] = 0
        counters["h3"] = 0
        prefix = f"{counters['h1'] - 1} "
        size = 14
    elif level == 2:
        counters["h2"] += 1
        counters["h3"] = 0
        prefix = f"{counters['h1'] - 1}.{counters['h2']} "
        size = 12
    else:
        counters["h3"] += 1
        prefix = f"{counters['h1'] - 1}.{counters['h2']}.{counters['h3']} "
        size = 11

    run = p.add_run(prefix + text)
    set_run_font(run, size_pt=size, bold=True)
    return p


def add_plain_heading(doc: Document, text: str, size: float = 14):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size_pt=size, bold=True)
    return p


def add_table(doc: Document, label: str):
    info = TABLE_DATA[label]
    add_plain_paragraph(doc, info["title"], align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, size=10)
    rows = len(info["rows"])
    cols = len(info["rows"][0])
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_idx, row in enumerate(info["rows"]):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = text
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    set_run_font(run, size_pt=9.5, bold=(r_idx == 0))
            set_cell_border(cell)
    if info["note"]:
        add_plain_paragraph(doc, info["note"], align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False, size=9)


def build_document():
    bib_map = parse_bib_map()
    refs = parse_references()
    blocks = parse_blocks(read_sections(), bib_map)
    extract_figure_asset()

    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Playwright 测试隐性误修复研究")
    set_run_font(r, size_pt=16, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Research on silent false healing in Playwright test repair")
    set_run_font(r, size_pt=12)

    for line in [
        "高东彪",
        "（软视软件（苏州）有限公司，江苏 苏州 215000）",
        "GAO Dongbiao",
        "(Zoom Video Communications (Suzhou), Inc., Suzhou 215000, China)",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run_font(r, size_pt=10.5)

    add_plain_paragraph(doc, "作者简介：高东彪（1986—），男，工程师，学士；研究方向：软件测试。", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    add_plain_paragraph(doc, "*通信作者：高东彪（1986—），男，工程师，学士；研究方向：软件测试。", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    add_plain_paragraph(doc, "摘 要：针对 Playwright 自动化测试中定位器易失效的问题，文章研究大语言模型驱动测试自愈的隐性误修复风险。该研究以 ReproBreak 数据集 koenig 项目的真实失效用例为对象，构建 HealReact 静态基底，并在选择器解析层面分析修复行为。结果表明，当目标元素缺失时，朴素修复器易生成指向错误元素的可解析选择器，而静态基底可覆盖多数真实失效目标。文章认为，持续集成场景应引入行为重放验证机制，以降低误修复风险。", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    add_plain_paragraph(doc, "关键词：测试自愈；大语言模型；Playwright；React；隐性误修复", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    add_plain_paragraph(doc, "中图分类号：TP311.5    文献标识码：A", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    add_plain_paragraph(doc, "Abstract: This paper investigates the risk of silent false healing in large language model (LLM)-driven self-healing for Playwright tests with brittle locators. Using real break cases from the koenig project in the ReproBreak dataset, the study builds HealReact, an abstract syntax tree-based static substrate, and analyzes repair behavior at the selector-resolution level. The results show that, when the target element is missing, a naive repairer easily generates resolver-valid selectors for wrong elements, while the static substrate covers most real break targets. The study indicates that continuous integration scenarios require a behavioral replay verification mechanism to reduce false-healing risk.", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    add_plain_paragraph(doc, "Key words: self-healing test; large language model; Playwright; React; silent false healing", align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)

    counters = {"h1": 0, "h2": 0, "h3": 0}
    for block in blocks:
        block_type = block["type"]
        if block_type == "heading1":
            add_heading(doc, block["text"], 1, counters)
        elif block_type == "heading2":
            add_heading(doc, block["text"], 2, counters)
        elif block_type == "heading3":
            add_heading(doc, block["text"], 3, counters)
        elif block_type == "paragraph":
            add_plain_paragraph(doc, block["text"])
        elif block_type == "bullet":
            add_plain_paragraph(doc, "• " + block["text"], first_line=False)
        elif block_type == "numbered":
            add_plain_paragraph(doc, block["text"], first_line=False)
        elif block_type == "figure":
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p_img.add_run()
            run.add_picture(str(FIGURE_IMAGE), width=Cm(14.5))
            add_plain_paragraph(doc, "图1 " + block["caption"], align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, size=9.5)
        elif block_type == "table":
            add_table(doc, block["label"])

    add_plain_heading(doc, "参考文献", size=14)
    for ref in refs:
        add_plain_paragraph(doc, ref, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False, size=10)

    doc.save(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()
